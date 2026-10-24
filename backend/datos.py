"""Datos ficticios de la demo: semilla del notebook, órdenes iniciales, reglas del agente y Excel.

El Excel usa exactamente el mismo formato que genera la celda 2.4 del notebook
(hojas TENANT, CONFIGURACION, COSTOS_CREDITOS, SUCURSALES, ..., EVENTOS_QUIEBRE).
"""
import base64
import copy
import io
import json
import os
from collections import OrderedDict

_RUTA = os.path.join(os.path.dirname(__file__), "semilla.json")
with open(_RUTA, encoding="utf-8") as _f:
    _SEMILLA = json.load(_f)

TABLAS = ["SUCURSALES", "ALMACENES", "USUARIOS", "PROVEEDORES", "PRODUCTOS", "STOCK",
          "VENTAS_DIARIAS", "COMPRAS_HISTORICAS", "OC_EN_TRANSITO", "EVENTOS_QUIEBRE"]
CAMPOS_LISTA = {"USUARIOS": ["permisos", "sucursales"]}
CONDICION_PAGO = {"PRV-01": "Contado", "PRV-02": "Crédito 30 días", "PRV-03": "Contado", "PRV-04": "Crédito 15 días"}


def semilla() -> dict:
    return copy.deepcopy(_SEMILLA)


# ─────────────────────────────────────────────────────────────── órdenes de compra iniciales
def _costo_ultima_compra(t, sku):
    compras = sorted((c for c in t["COMPRAS_HISTORICAS"] if c["sku"] == sku), key=lambda c: c["fecha"])
    return compras[-1]["costo_unitario"] if compras else None


def _linea(t, sku, almacen_id, cantidad, recibida=0):
    alm = next(a for a in t["ALMACENES"] if a["almacen_id"] == almacen_id)
    return {"sku": sku, "almacen_id": almacen_id, "sucursal_id": alm["sucursal_id"], "cantidad": cantidad,
            "recibida": recibida, "costo_unitario": _costo_ultima_compra(t, sku), "sugerida": None,
            "detalle_sugerencia": None, "evento_ids": []}


def _orden(numero, proveedor_id, estado, creada_por, creada_en, entrega, lineas, enviada=None):
    historial = [{"ts": creada_en, "accion": "Creada", "actor": creada_por}]
    if estado in ("APROBADA", "RECIBIDA_PARCIAL", "RECIBIDA"):
        historial.append({"ts": creada_en, "accion": "Aprobada", "actor": "U01"})
    if enviada:
        historial.append({"ts": enviada, "accion": "Enviada al proveedor (PIN)", "actor": "U01"})
    if estado == "PENDIENTE_APROBACION":
        historial.append({"ts": creada_en, "accion": "Enviada a aprobación: supera el tope", "actor": creada_por})
    return {"numero": numero, "proveedor_id": proveedor_id, "estado": estado, "origen": "MANUAL", "corrida_id": None,
            "propuesta_id": None, "creada_por": creada_por, "creada_en": creada_en, "fecha_entrega": entrega,
            "condicion_pago": CONDICION_PAGO.get(proveedor_id, "Contado"), "lineas": lineas,
            "enviada_al_proveedor": {"ts": enviada, "por": "U01", "canal": "email"} if enviada else None,
            "sin_proveedor_asignado": False, "historial": historial}


def ordenes_iniciales(t) -> list:
    """Órdenes en tránsito del Excel (OC_EN_TRANSITO) y, con los datos de ejemplo, 3 órdenes más para la tabla."""
    agrupadas = OrderedDict()
    for fila in t["OC_EN_TRANSITO"]:
        agrupadas.setdefault(fila["oc_id"], []).append(fila)
    ocs = []
    for oc_id, filas in agrupadas.items():
        f0 = filas[0]
        parcial = f0["estado"] == "RECIBIDA_PARCIAL"
        lineas = [_linea(t, f["sku"], f["almacen_id"], f["cantidad_pendiente"] + (4 if parcial else 0), 4 if parcial else 0)
                  for f in filas]
        ocs.append(_orden(oc_id, f0["proveedor_id"], "RECIBIDA_PARCIAL" if parcial else "APROBADA", "U01",
                          f0["fecha_envio"] + "T10:00:00", f0["fecha_estimada"], lineas, enviada=f0["fecha_envio"] + "T10:30:00"))
    skus = {p["sku"] for p in t["PRODUCTOS"]}
    if {"COL-PER-ROS", "ANI-PLA-AJU", "SET-NOV-PLA"} <= skus and "OC-0143" not in agrupadas:
        r = _orden("OC-0140", "PRV-01", "RECIBIDA", "U01", "2026-10-03T11:00:00", "2026-10-08",
                   [_linea(t, "COL-PER-ROS", "ALM-01", 20, 20)], enviada="2026-10-03T11:20:00")
        a = _orden("OC-0143", "PRV-04", "APROBADA", "U01", "2026-10-13T16:00:00", "2026-10-20",
                   [_linea(t, "ANI-PLA-AJU", "ALM-02", 30)], enviada="2026-10-13T16:10:00")
        p = _orden("OC-0146", "PRV-01", "PENDIENTE_APROBACION", "U01", "2026-10-21T12:00:00", "2026-10-28",
                   [_linea(t, "SET-NOV-PLA", "ALM-01", 10)])
        ocs = [r, ocs[0], a] + ocs[1:] + [p]
    return ocs


# ─────────────────────────────────────────────────────────────── reglas del agente (corpus)
REGLAS_AGENTE = [
    {"id": "AB-01", "titulo": "Qué hace y qué no hace el agente",
     "texto": "Detecta qué reponer, calcula la cantidad sugerida, busca el proveedor de la última compra, agrupa por proveedor y deja Borradores de Orden de Compra con origen 'generada por agente'. Revisar, ajustar, aprobar, despachar y recibir siguen siendo humanos. El agente nunca aprueba, despacha, recibe ni paga."},
    {"id": "AB-02", "titulo": "Cuándo se ejecuta",
     "texto": "Tres mecanismos: eventos de quiebre de core-stock que solo se acumulan en una cola; una corrida programada diaria a la hora elegida por el negocio; y la tarjeta 'Sugerir Recompra' a demanda. Una corrida programada sin eventos nuevos se omite sin consumir créditos. Escribir 'haz una orden de compra' en el chat no dispara al agente."},
    {"id": "AB-03", "titulo": "Cómo se calcula la cantidad sugerida",
     "texto": "Velocidad diaria = unidades vendidas en la ventana (30 días) / 30. Reposición sugerida = velocidad × días de cobertura (15) − stock disponible, hacia arriba y nunca negativa. Faltante al mínimo = mínimo − stock. Necesidad = el mayor de ambos. Cantidad sugerida = necesidad − unidades en tránsito. No es un lote económico (EOQ) ni una cantidad óptima."},
    {"id": "AB-04", "titulo": "Órdenes en tránsito",
     "texto": "Las unidades de órdenes ya enviadas al proveedor y aún no recibidas se descuentan. Si cubren la necesidad, la línea se excluye con el motivo 'cubierto por OC en tránsito'."},
    {"id": "AB-05", "titulo": "A qué proveedor se le pide",
     "texto": "Cada producto va al proveedor de su última compra. Un borrador por proveedor. Un producto sin compras previas va a un borrador 'sin proveedor asignado' que no se puede enviar a aprobación hasta elegir proveedor."},
    {"id": "AB-06", "titulo": "Alcance multi-sucursal",
     "texto": "Consolidado (una orden por proveedor con todas las sucursales), por separado (una por sucursal y proveedor) o solo algunas sucursales."},
    {"id": "AB-07", "titulo": "PIN: proponer no requiere PIN, ejecutar sí",
     "texto": "Crear un borrador no exige PIN. Enviar la orden al proveedor sí exige el PIN del Administrador (o el código de contingencia por WhatsApp) y se hace en Compras. El PIN nunca se escribe en el chat."},
    {"id": "AB-08", "titulo": "Créditos de IA",
     "texto": "Cada corrida y cada consulta descuentan créditos. Sin saldo suficiente la corrida no se inicia. El chat se pausa al llegar a la reserva para agentes. Un error del modelo o una corrida omitida nunca se cobran."},
    {"id": "AB-09", "titulo": "Requisitos",
     "texto": "Add-on IA Multi-Agente activo y un plan con el Módulo de Órdenes de Compra. El agente nace apagado y se activa en Configuración › IA."},
    {"id": "AB-10", "titulo": "Auditoría",
     "texto": "Toda operación de una corrida se registra con el agente como actor y el id de corrida. Cuando una persona aprueba o envía con PIN, se registra un segundo evento con su identidad."},
    {"id": "AB-11", "titulo": "Permisos y visibilidad",
     "texto": "El agente responde con los permisos de quien pregunta: sin 'ver costos' no muestra costos y solo muestra el stock de las sucursales asignadas."},
    {"id": "AB-12", "titulo": "Aprobación por monto",
     "texto": "Al enviar a aprobación, una orden por encima del tope queda Pendiente de aprobación; por debajo pasa directo a Aprobada. Nadie aprueba lo que creó; las órdenes del agente sí las puede aprobar cualquier usuario con permiso."},
]


# ─────────────────────────────────────────────────────────────── Excel
def _openpyxl():
    import openpyxl  # dependencia solo para importar/exportar
    return openpyxl


def leer_excel_b64(contenido_b64: str) -> dict:
    """Lee un .xlsx con el formato del notebook y devuelve las tablas. Lanza ValueError con un mensaje claro."""
    openpyxl = _openpyxl()
    try:
        libro = openpyxl.load_workbook(io.BytesIO(base64.b64decode(contenido_b64)), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"No se pudo abrir el Excel: {type(exc).__name__}")

    def filas(hoja):
        if hoja not in libro.sheetnames:
            raise ValueError(f"Falta la hoja {hoja}")
        it = libro[hoja].iter_rows(values_only=True)
        cab = [str(c).strip() if c is not None else "" for c in next(it)]
        return [{k: v for k, v in zip(cab, fila) if k} for fila in it if any(x is not None for x in fila)]

    t = semilla()
    t["TENANT"] = {r["clave"]: json.loads(r["valor"]) for r in filas("TENANT")}
    t["CONFIGURACION"] = {r["clave"]: json.loads(r["valor"]) for r in filas("CONFIGURACION")}
    t["COSTOS_CREDITOS"] = {r["operacion"]: int(r["creditos"]) for r in filas("COSTOS_CREDITOS")}
    for nombre in TABLAS:
        registros = []
        for f in filas(nombre):
            f = {k: (v.isoformat()[:10] if hasattr(v, "isoformat") and k in ("fecha", "fecha_envio", "fecha_estimada") else v)
                 for k, v in f.items()}
            for campo in CAMPOS_LISTA.get(nombre, []):
                f[campo] = [x.strip() for x in str(f.get(campo) or "").split(",") if x.strip()]
            for campo in ("ruc",):
                if campo in f and f[campo] is not None:
                    f[campo] = str(f[campo])
            registros.append(f)
        t[nombre] = registros
    if len(t["VENTAS_DIARIAS"]) > 20000 or len(t["PRODUCTOS"]) > 500:
        raise ValueError("La demo admite hasta 500 productos y 20.000 filas de ventas")
    return t


def validar(t: dict) -> list:
    errores = []
    skus = {p["sku"] for p in t["PRODUCTOS"]}
    almacenes = {a["almacen_id"] for a in t["ALMACENES"]}
    proveedores = {p["proveedor_id"] for p in t["PROVEEDORES"]}
    sucursales = {s["sucursal_id"] for s in t["SUCURSALES"]}
    claves = set()
    for s in t["STOCK"]:
        if s["sku"] not in skus: errores.append(f"STOCK: SKU desconocido {s['sku']}")
        if s["almacen_id"] not in almacenes: errores.append(f"STOCK: almacén desconocido {s['almacen_id']}")
        if (s["sku"], s["almacen_id"]) in claves: errores.append(f"STOCK: fila duplicada {s['sku']}/{s['almacen_id']}")
        claves.add((s["sku"], s["almacen_id"]))
    for c in t["COMPRAS_HISTORICAS"] + t["OC_EN_TRANSITO"]:
        if c["proveedor_id"] not in proveedores: errores.append(f"Proveedor desconocido {c['proveedor_id']}")
        if c["sku"] not in skus: errores.append(f"SKU desconocido {c['sku']}")
    for e in t["EVENTOS_QUIEBRE"]:
        if (e["sku"], e["almacen_id"]) not in claves: errores.append(f"Evento {e['evento_id']} sin fila de stock")
    for a in t["ALMACENES"]:
        if a["sucursal_id"] not in sucursales: errores.append(f"Almacén {a['almacen_id']} con sucursal desconocida")
    if t["CONFIGURACION"].get("configurado_por") not in {u["usuario_id"] for u in t["USUARIOS"]}:
        errores.append("CONFIGURACION: configurado_por no existe en USUARIOS")
    return errores[:15]


def exportar_excel(S) -> bytes:
    """Exporta el estado actual de la demo (con las ventas simuladas) en el formato del notebook."""
    openpyxl = _openpyxl()
    libro = openpyxl.Workbook()
    libro.remove(libro.active)

    def hoja(nombre, filas, columnas=None):
        ws = libro.create_sheet(nombre)
        columnas = columnas or (list(filas[0].keys()) if filas else [])
        ws.append(columnas)
        for f in filas:
            ws.append([",".join(f[c]) if isinstance(f.get(c), list) else f.get(c) for c in columnas])

    d = S["datos"]
    hoja("TENANT", [{"clave": k, "valor": json.dumps(v, ensure_ascii=False)} for k, v in S["tenant"].items()])
    hoja("CONFIGURACION", [{"clave": k, "valor": json.dumps(v, ensure_ascii=False)} for k, v in S["config"].items()])
    hoja("COSTOS_CREDITOS", [{"operacion": k, "creditos": v} for k, v in S["costos_creditos"].items()])
    for nombre in ("SUCURSALES", "ALMACENES", "USUARIOS", "PROVEEDORES", "PRODUCTOS", "STOCK", "VENTAS_DIARIAS", "COMPRAS_HISTORICAS"):
        hoja(nombre, d[nombre])
    transito = [{"oc_id": o["numero"], "proveedor_id": o["proveedor_id"],
                 "estado": "RECIBIDA_PARCIAL" if o["estado"] == "RECIBIDA_PARCIAL" else "ENVIADA",
                 "fecha_envio": (o["enviada_al_proveedor"]["ts"] or "")[:10], "fecha_estimada": o["fecha_entrega"],
                 "sku": l["sku"], "almacen_id": l["almacen_id"], "cantidad_pendiente": l["cantidad"] - l["recibida"]}
                for o in S["ocs"] if o["enviada_al_proveedor"] and o["estado"] in ("APROBADA", "RECIBIDA_PARCIAL")
                for l in o["lineas"] if l["cantidad"] > l["recibida"]]
    hoja("OC_EN_TRANSITO", transito, ["oc_id", "proveedor_id", "estado", "fecha_envio", "fecha_estimada", "sku", "almacen_id", "cantidad_pendiente"])
    hoja("EVENTOS_QUIEBRE", S["eventos"], ["evento_id", "fecha_hora", "sku", "almacen_id", "tipo", "procesado"])
    buf = io.BytesIO()
    libro.save(buf)
    return buf.getvalue()
